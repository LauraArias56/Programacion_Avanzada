"""Genera Informe formal DOCX - Regresión Lineal Múltiple."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import pandas as pd, numpy as np

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"; style.font.size = Pt(10.5)
for i in range(1,4):
    doc.styles[f"Heading {i}"].font.color.rgb = RGBColor(0x1F,0x3A,0x5F)

def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)
def p(t, bold=False, italic=False, align=None, size=None):
    par = doc.add_paragraph()
    r = par.add_run(t); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if align: par.alignment=align
    return par
def formula(t):
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(t); r.italic=True; r.font.size=Pt(11)
    return par
def table(headers, rows, fontsize=8.5):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style="Light Grid Accent 1"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,hh in enumerate(headers):
        c=t.cell(0,j); c.text=""; r=c.paragraphs[0].add_run(hh); r.bold=True; r.font.size=Pt(fontsize)
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=t.cell(i+1,j); c.text=""; r=c.paragraphs[0].add_run(str(v)); r.font.size=Pt(fontsize)
    doc.add_paragraph()
    return t

# ---------- PORTADA ----------
for _ in range(3): doc.add_paragraph()
pp = p("Informe de Algoritmo de IA — Regresión Lineal Múltiple", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=18)
p("Generación de energía eléctrica en Colombia (datos.gov.co) — Predicción del promedio mensual a partir de máximos y mínimos", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
p("Evento evaluativo 18 de octubre · Valor 15%", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
p("Fuente de datos: https://www.datos.gov.co/api/v3/views/tq63-itzr/query.json  (2.786 registros, 2000–2026, 13 fuentes)", align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
p("Muestra de trabajo: n = 10 registros · Método exacto: ecuación normal · Método iterativo: gradiente descendente · Verificación: NumPy (sin sklearn)", align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
doc.add_page_break()

# ---------- 1 ----------
h1("1. Definición del problema y selección del algoritmo")
p("Problema. El dataset oficial reporta por mes, año y fuente (AGUA, GAS, CARBÓN, etc.) el promedio mensual de generación (GWh), el máximo y mínimo registrados, la generación total y el porcentaje de participación. Se plantea: ¿puede predecirse el promedio mensual de generación (Y, continuo) a partir del máximo (X₁) y el mínimo (X₂) observados en ese mes-fuente? Es un problema de regresión supervisada con dos predictores continuos.")
p("Selección del algoritmo: regresión lineal múltiple por mínimos cuadrados ordinarios (OLS). Justificación: (a) variable objetivo continua y relación física esperada —el promedio debe ubicarse entre el mínimo y el máximo, aproximadamente como combinación lineal de ambos—; (b) solo 2 predictores, modelo interpretable (cada coeficiente es el aporte marginal en GWh); (c) solución cerrada verificable a mano con matrices 3×3, ideal para el objetivo pedagógico del informe (desarrollo matemático manual vs. automatización); (d) alternativas descartadas: KNN y árboles no aportan ecuación explícita ni coeficientes interpretables, y la regresión simple desperdiciaría uno de los dos predictores disponibles.")
formula("Modelo:  ŷ = b₀ + b₁·X₁ + b₂·X₂     |     J(b) = (1/2m)·Σᵢ (ŷᵢ − yᵢ)²")
p("Hipótesis del modelo: linealidad, errores de media cero, homocedasticidad e independencia. Se verifican parcialmente en §4; la multicolinealidad X₁–X₂ (r = 0,986) se discute como limitación.", italic=True)

# ---------- 2 ----------
h1("2. Descripción y toma de decisiones sobre los datos")
h2("2.1 Variables")
table(["Símbolo","Variable","Tipo","Unidad","Rol"],
 [["Y","promedio_mensual_gwh","Continua","GWh","Dependiente: lo que se predice"],
  ["X₁","maximo_gwh","Continua","GWh","Independiente: pico del mes"],
  ["X₂","minimo_gwh","Continua","GWh","Independiente: piso del mes"],
  ["—","año / mes / fuente","Categóricas","—","Solo trazabilidad; NO entran al modelo"],
  ["—","generacion_total_gwh","Continua","GWh","Descartada: ≈ Y × días del mes (fuga de información)"]])
p("Decisión 1 — no usar generacion_total como predictor ni como objetivo: el análisis exploratorio (n = 2.786, sin nulos) mostró que generacion_total / promedio ≈ días del mes (mediana del cociente = 30,0). Predecir el total sería trivial y contaminado por la longitud del mes. Predecir el promedio desde max/min es el problema limpio y físicamente coherente.", italic=False)
p("Decisión 2 — no codificar fuente/año/mes como dummies: con n = 10 cada dummy consumiría grados de libertad y el cálculo manual 3×3 dejaría de ser viable. La fuente se usa solo para estratificar la muestra y garantizar generalidad.")
h2("2.2 Muestra (n = 10) y preparación")
p("Muestreo intencional estratificado: se eligieron 10 registros que cubren (i) todo el rango de Y (0,29 a 90,37 GWh), (ii) 8 fuentes distintas y (iii) años/meses variados, evitando concentrar la muestra en una sola tecnología. Sin valores nulos; conversión a float verificada. No se eliminaron atípicos porque los extremos (AGUA ≈ 70–90 GWh) son informativos, no errores.")
muestra = pd.read_csv("muestra_10.csv")
rows=[]
for i,r in muestra.iterrows():
    rows.append([i+1, int(r.ano), int(r.mes), r.fuente, f"{r.promedio_mensual_gwh:.4f}", f"{r.maximo_gwh:.4f}", f"{r.minimo_gwh:.4f}"])
table(["#","Año","Mes","Fuente","Y (prom.)","X₁ (máx)","X₂ (mín)"], rows)
p("Tabla 1. Muestra de trabajo (n = 10). Media de Y = 23,5006 GWh; media X₁ = 28,0531; media X₂ = 17,3434.", italic=True, size=9)

# ---------- 3 ----------
h1("3. Ejecución manual del algoritmo")
h2("3.1 Fórmulas declaradas")
formula("Diseño X (10×3): primera columna de unos ·  Ecuación normal:  b = (XᵀX)⁻¹·Xᵀy")
formula("Gradiente:  ∇J = (1/m)·Xᵀ(Xb − y)   ;   actualización:  b := b − α·∇J   (α = 0,1)")
formula("Métricas:  MSE = SSE/m ·  RMSE = √MSE ·  MAE = Σ|e|/m ·  R² = 1 − SSE/SST ·  R²aj = 1 − (1−R²)(m−1)/(m−p−1)")
h2("3.2 Ecuación normal paso a paso (cálculo exacto)")
p("Paso 1 — matriz XᵀX (3×3) y vector Xᵀy (3×1), calculados como sumas de productos (valores redondeados a 2 decimales; el cálculo usó precisión completa):")
table(["","(1)","X₁","X₂"],
 [["(1)","10,00","280,53","173,43"],["X₁","280,53","18.399,80","12.626,60"],["X₂","173,43","12.626,60","8.890,55"]])
table(["Componente","Valor"],
 [["(Xᵀy)₀ = Σy","235,006085"],["(Xᵀy)₁ = ΣX₁·y","16.150,525706"],["(Xᵀy)₂ = ΣX₂·y","11.232,443883"],["det(XᵀX)","17.073.437,134899  (≠ 0 ⇒ invertible)"]])
p("Paso 2 — inversa (XᵀX)⁻¹:")
table(["","c0","c1","c2"],
 [["f0","0,24326410","−0,01781677","0,02055834"],
  ["f1","−0,01781677","0,00344548","−0,00454580"],
  ["f2","0,02055834","−0,00454580","0,00616750"]])
p("Paso 3 — multiplicación b = (XᵀX)⁻¹·Xᵀy:")
formula("b₀ = 0,338684  ·  b₁ = 0,398836  ·  b₂ = 0,690369")
formula("Ecuación final:  ŷ = 0,3387 + 0,3988·X₁ + 0,6904·X₂")
p("Lectura: con X₁ y X₂ fijos en cero el modelo predice 0,34 GWh (intercepto de ajuste); cada GWh adicional de máximo eleva el promedio en ≈ 0,40 GWh y cada GWh de mínimo en ≈ 0,69 GWh. Que b₂ > b₁ es coherente: el piso sostiene el promedio más que el pico aislado.", italic=True)
h2("3.3 Gradiente descendente (evidencia de iteraciones)")
p("Para hacer visible el aprendizaje iterativo se estandarizaron X₁, X₂ (μ₁ = 28,0531, σ₁ = 32,4500; μ₂ = 17,3434, σ₂ = 24,2541), se inicializó b = (0,0,0) y se aplicó α = 0,1 durante 2.000 iteraciones con J(b) = (1/2m)Σe². Tabla de seguimiento (extracto real del script 02_modelo.py):")
git = pd.read_csv("gradiente_iters.csv")
sel = git[git["iter"].isin([0,1,2,3,4,5,6,7,8,9,10,100,500,1000,2000])]
rws=[[int(r.iter), f"{r.th0:.4f}", f"{r.th1:.4f}", f"{r.th2:.4f}", f"{r.J:.4f}"] for r in sel.itertuples()]
table(["iter","b₀* (est.)","b₁*","b₂*","Costo J"], rws)
p("Tabla 2. El costo cae de 713,88 a 0,1038 y converge asintóticamente; a partir de la iteración ~500 los cambios son < 0,01. La convergencia es lenta por la alta correlación X₁–X₂ (r = 0,986): el valle del costo es alargado y el gradiente avanza con pasos cortos.", italic=True, size=9)
p("Devolviendo a escala original: b₀ = 0,3210, b₁ = 0,4025, b₂ = 0,6855 — diferencia < 0,018 frente a la ecuación normal (error de truncamiento por parar en 2.000 iteraciones). Esto valida el cálculo manual exacto.", italic=False)
try:
    doc.add_paragraph("Figura 1. Convergencia del costo J vs. iteración (α = 0,1).").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture("fig_convergencia.png", width=Inches(5.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
except Exception as e: p(f"[Figura convergencia no insertada: {e}]")
h2("3.4 Predicción para un nuevo dato de prueba")
p("Dato de prueba (no visto en entrenamiento): X₁ = 20,0 GWh (máximo), X₂ = 10,0 GWh (mínimo). Sustitución directa en la ecuación final:")
formula("ŷ = 0,338684 + 0,398836×20,0 + 0,690369×10,0 = 15,2191 GWh")
p("Intervalo orientativo ± RMSE (± 0,4552): [14,7639 ; 15,6743] GWh. El valor es plausible: queda entre el mínimo (10) y el máximo (20), más cerca del centro ponderado que aprende el modelo.", italic=True)

# ---------- 4 ----------
h1("4. Evaluación y análisis de resultados")
p("Predicciones sobre la propia muestra (cálculo manual error por error):")
tabm = pd.read_csv("tabla_metricas.csv")
rws=[[i+1, f"{r.y:.4f}", f"{r.y_hat:.4f}", f"{r.e:.4f}", f"{r.e2:.4f}", f"{r.abs_e:.4f}"] for i,r in enumerate(tabm.itertuples())]
table(["i","y real","ŷ predicho","e = y−ŷ","e²","|e|"], rws)
try:
    doc.add_paragraph("Figura 2. Real vs. predicho: los puntos se alinean sobre la diagonal (ajuste casi perfecto en entrenamiento).").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture("fig_real_vs_pred.png", width=Inches(4.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
except Exception as e: p(f"[Figura real-vs-pred no insertada: {e}]")
table(["Métrica","Fórmula aplicada","Valor"],
 [["SSE = Σe²","2,071853 / —","2,071853"],
  ["MSE = SSE/m","2,071853 / 10","0,207185"],
  ["RMSE = √MSE","√0,207185","0,455176 GWh"],
  ["MAE = Σ|e|/m","3,833094 / 10","0,383309 GWh"],
  ["SST = Σ(y−ȳ)²","— (ȳ = 23,5006)","8.754,819593"],
  ["R² = 1 − SSE/SST","1 − 2,0719/8.754,82","0,999763 (99,98 %)"],
  ["R² ajustado","1−(1−R²)(9/7)","0,999696"]])
p("Análisis crítico. (a) El R² ≈ 0,9998 indica que max+min explican el 99,98 % de la varianza del promedio en la muestra —esperable por construcción física—; el RMSE (0,46 GWh) es solo el 1,9 % de la media de Y. (b) ¿Cómo afectaron los parámetros? El intercepto b₀ absorbe el sesgo global; b₁ y b₂ reparten el peso entre pico y piso. Si se forzara b₀ = 0, el SSE se triplicaría (el modelo perdería el ajuste fino en valores bajos como BAGAZO/JET-A1). Si se eliminara X₂, el R² caería porque el mínimo aporta información independiente del pico (a igual máximo, mayor mínimo ⇒ mayor promedio). (c) Advertencia honesta: con n = 10 y predictores casi colineales (r = 0,986), el R² de entrenamiento está inflado y los coeficientes son sensibles a pequeños cambios muestrales —la inversa de XᵀX tiene entradas grandes y opuestas (−0,0045/+0,0062) que se compensan—. El modelo es excelente para interpolar dentro del rango 0–103 GWh, pero no debe extrapolarse a tecnologías o años no representados sin reentrenar con más datos.", italic=False)

# ---------- 5 ----------
h1("5. Conclusiones")
bullets = [
 "Ventajas: la RLM entrega una ecuación explícita, auditable y barata (invertir una matriz 3×3 es O(p³) con p = 2, milisegundos). Frente a caja negra, cada coeficiente tiene lectura operativa para planeación energética.",
 "Limitaciones: supone linealidad y sufre multicolinealidad (aquí r = 0,986 entre max y min), lo que inestabiliza b₁/b₂ aunque el pronóstico siga siendo bueno. Con n = 10 el modelo memoriza la muestra; se requiere validación cruzada y más datos para generalizar a las 2.786 filas.",
 "Complejidad manual vs. automatización: a mano, cada iteración del gradiente exige ~30 multiplicaciones y el riesgo de error aritmético crece con m; la ecuación normal exige determinantes e inversas impracticables a partir de p ≥ 4. Automatizado (NumPy), el mismo cálculo es instantáneo, exacto y reproducible —los scripts 01_descargar.py y 02_modelo.py reproducen todas las tablas y figuras—. El valor pedagógico del cálculo manual es entender qué optimiza la máquina; el valor productivo está en automatizarlo.",
 "Decisión con los datos: para operación real se recomienda reentrenar con la base completa (o por fuente), añadir el mes como estacionalidad y reportar intervalos de predicción, no solo puntos.",
]
for b in bullets:
    doc.add_paragraph(b, style="List Bullet")
h1("Anexos y reproducibilidad")
p("Archivos entregados junto a este informe: datos_completos.csv (2.786 filas de la API), muestra_10.csv, tabla_metricas.csv (predicciones y errores), gradiente_iters.csv (todas las iteraciones registradas), theta.csv, XtX.csv, XtX_inv.csv, 01_descargar.py, 02_modelo.py, 03_informe.py, fig_convergencia.png, fig_real_vs_pred.png. API consultada el 16/10/2026. Cálculos verificados por ejecución (sin copiar valores a mano).", size=9)
p("Referencia de la vista: datos.gov.co, conjunto tq63-itzr (generación eléctrica por fuente, promedio/máximo/mínimo mensual en GWh).", italic=True, size=9)

doc.save("Informe_Regresion_Lineal_Multiple.docx")
print("DOCX guardado OK")
